"""import_category.py's one pure piece: reading a merchandising class export
into {id, category, subcategory} rows.

The export's printed header names only 5 of the 8 actual columns per row
(confirmed by inspecting data/all_products_class.xlsx directly), so
read_classes() reads by position, not by header name -- the header row is
skipped, not parsed. Price Class (col3) is confirmed to be the real main
category field despite the name, not a pricing tier. The safety property
that matters: a row with a blank Price Class must be dropped rather than
turned into a row that would clear an existing category to blank.
"""
import openpyxl
import pytest

from import_category import read_classes


@pytest.fixture
def class_xlsx(tmp_path):
    path = tmp_path / "class.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    # The real export's header names only 5 columns; every data row still
    # carries all 8 values. read_classes ignores this row's text entirely.
    ws.append(["Item ", "Description", "Sales Category", "Price Class", "OTP Level", None, None, None])
    ws.append([1358, "100 GRAND CHOCOLATE BAR 36CT", 6, "FOOD", 129, "CHOCOLATE", 0, "N/A"])
    ws.append([1203, "12PK CANADA DRY", 5, "DRINK", 24, "12PK COKE", 0, "N/A"])
    ws.append([9999, "NO CATEGORY ROW", None, "", None, "", 0, "N/A"])  # blank category -- dropped
    ws.append([None, "NO ITEM NUMBER", 6, "FOOD", 129, "CHOCOLATE", 0, "N/A"])  # no id -- dropped
    wb.save(path)
    return path


class TestReadClasses:
    def test_reads_id_category_subcategory_by_position(self, class_xlsx):
        rows = read_classes(class_xlsx)
        assert {"id": "1358", "category": "FOOD", "subcategory": "CHOCOLATE"} in rows
        assert {"id": "1203", "category": "DRINK", "subcategory": "12PK COKE"} in rows

    def test_rows_with_a_blank_category_are_dropped(self, class_xlsx):
        rows = read_classes(class_xlsx)
        assert all(r["id"] != "9999" for r in rows)

    def test_rows_without_an_item_number_are_dropped(self, class_xlsx):
        rows = read_classes(class_xlsx)
        assert all(r["id"] for r in rows)
        assert len(rows) == 2

    def test_ids_are_strings(self, class_xlsx):
        rows = read_classes(class_xlsx)
        assert all(isinstance(r["id"], str) for r in rows)

    def test_missing_subcategory_is_none_not_a_placeholder(self, tmp_path):
        path = tmp_path / "no_sub.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Item ", "Description", "Sales Category", "Price Class", "OTP Level"])
        ws.append([100, "SOMETHING", 6, "FOOD", None])
        wb.save(path)

        rows = read_classes(path)
        assert rows == [{"id": "100", "category": "FOOD", "subcategory": None}]

    def test_short_rows_are_dropped_without_crashing(self, tmp_path):
        # openpyxl trims trailing all-empty cells -- a row shorter than the
        # category column's position must not raise an IndexError.
        path = tmp_path / "short_rows.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Item ", "Description", "Sales Category", "Price Class", "OTP Level"])
        ws.append([100, "TRUNCATED"])
        wb.save(path)

        rows = read_classes(path)
        assert rows == []
