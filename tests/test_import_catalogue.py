"""import_catalogue.py's two pure pieces: reading a catalogue export into
plain rows, and diffing against what the products table already has. The
safety property that matters is in not_yet_in_table -- it must never hand
back a row whose id already exists, since that row is what gets pushed as a
status='not_sourced' stub, and overwriting an already-sourced product with
one would erase real pipeline output.
"""
import openpyxl
import pytest

from import_catalogue import not_yet_in_table, read_catalogue


@pytest.fixture
def catalogue_xlsx(tmp_path):
    path = tmp_path / "catalogue.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Item Number", "Pack", "Size / UOM", "Description", "Retail", "Price"])
    ws.append([100, 12, "12 OZ", "FIRST PRODUCT", 1.99, 20.0])
    ws.append([200, 1, "1 LITER", "SECOND PRODUCT", 2.99, 30.0])
    ws.append([None, 1, "", "*B/O*", 0, 0])  # no item number -- must be dropped
    wb.save(path)
    return path


class TestReadCatalogue:
    def test_reads_id_and_name(self, catalogue_xlsx):
        rows = read_catalogue(catalogue_xlsx)
        assert {"id": "100", "name": "FIRST PRODUCT"} in rows
        assert {"id": "200", "name": "SECOND PRODUCT"} in rows

    def test_rows_without_an_item_number_are_dropped(self, catalogue_xlsx):
        rows = read_catalogue(catalogue_xlsx)
        assert all(r["id"] for r in rows)
        assert len(rows) == 2

    def test_ids_are_strings(self, catalogue_xlsx):
        # openpyxl hands back a numeric cell as an int/float; the products
        # table's id column is text, and every other id in this codebase is
        # a string ("100060", not 100060).
        rows = read_catalogue(catalogue_xlsx)
        assert all(isinstance(r["id"], str) for r in rows)

    def test_missing_description_gets_a_placeholder(self, tmp_path):
        path = tmp_path / "no_name.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Item Number", "Description"])
        ws.append([100, None])
        wb.save(path)

        rows = read_catalogue(path)
        assert rows == [{"id": "100", "name": "(no description)"}]

    def test_reads_the_merged_master_sheet_header_shape(self, tmp_path):
        # build_master_sheet.py's own output uses "id"/"name", not
        # "Item Number"/"Description" -- import_catalogue.py's new default
        # (data/master_products.xlsx) must be readable, not just the
        # original per-batch export shape.
        path = tmp_path / "master.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["id", "name", "pack", "size", "retail", "price",
                   "category", "subcategory", "otp_level",
                   "in_catalogue_file", "in_class_file"])
        ws.append([100, "MERGED PRODUCT", 12, "12 OZ", 1.99, 20.0,
                   "FOOD", "SNACKS", None, True, True])
        wb.save(path)

        rows = read_catalogue(path)
        assert rows == [{"id": "100", "name": "MERGED PRODUCT"}]


class TestNotYetInTable:
    def test_keeps_rows_whose_id_is_absent(self):
        catalogue = [{"id": "100", "name": "A"}, {"id": "200", "name": "B"}]
        assert not_yet_in_table(catalogue, have=set()) == catalogue

    def test_drops_rows_already_present(self):
        catalogue = [{"id": "100", "name": "A"}, {"id": "200", "name": "B"}]
        missing = not_yet_in_table(catalogue, have={"100"})
        assert missing == [{"id": "200", "name": "B"}]

    def test_never_returns_an_id_that_is_in_have(self):
        """The safety property: whatever comes back must not overwrite an
        already-sourced product, no matter what `have` contains."""
        catalogue = [{"id": str(i), "name": f"item {i}"} for i in range(50)]
        have = {str(i) for i in range(0, 50, 2)}  # every even id already exists

        missing = not_yet_in_table(catalogue, have)

        assert not any(r["id"] in have for r in missing)

    def test_everything_present_returns_empty(self):
        catalogue = [{"id": "100", "name": "A"}]
        assert not_yet_in_table(catalogue, have={"100"}) == []
